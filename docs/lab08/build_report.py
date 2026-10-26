"""Builds the report of laboratory work 8 (RWPSF 3305) from the project sources and docs/lab08.

Run from any directory:  python docs/lab08/build_report.py
Listings are read from the real source files, figures are the captures in docs/lab08/img,
the numbers in the tables are the ones measured in docs/lab08/out.
"""
import json
import re
import sys
from pathlib import Path

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

sys.path.insert(0, r'C:\Users\kuatw\.claude\skills\lab-title-page')
from title_page import new_report  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
LAB = ROOT / 'docs' / 'lab08'
IMG, OUT = LAB / 'img', LAB / 'out'
JAVA = ROOT / 'src' / 'main' / 'java' / 'kz' / 'iitu' / 'library'
BRANCH = 'https://github.com/AkAyu44/library-portal/tree/lab08'
COMMIT = 'https://github.com/AkAyu44/library-portal/commit/dc290ed'

doc = new_report(
    discipline='RWPSF 3305 – Web Application Development with Spring Framework',
    lab='Laboratory Work 8. The data access layer',
    subtitle='Connecting PostgreSQL, entities and associations (@OneToMany, @ManyToMany) · '
             'Spring Data JPA repositories, derived and JPQL queries · pagination and sorting',
    date='26 October 2026',
    extras=[('Lecturer', 'Yemberdiyeva Aknur, Senior Lector'),
            ('Individual variant', 'No. 7 — Window and keyset scrolling'),
            ('Repository', 'github.com/AkAyu44/library-portal, branch lab08')],
)
doc.core_properties.author = 'Kuat Bektursyn'
doc.core_properties.last_modified_by = 'Kuat Bektursyn'
doc.core_properties.title = 'Laboratory Work 8. The data access layer'

FIG, TAB, LST = {}, {}, {}


# ---------------------------------------------------------------- helpers
def shade(cell, hex_fill):
    pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    pr.append(shd)


def h1(text):
    doc.add_heading(text, 1)


def h2(text):
    doc.add_heading(text, 2)


def para(text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
    p.add_run(text)
    p.paragraph_format.space_after = Pt(6)
    return p


def rich(*parts):
    """parts: plain str, or ('b' | 'i' | 'c', text) for bold, italic, code."""
    p = doc.add_paragraph()
    for part in parts:
        if isinstance(part, str):
            p.add_run(part)
            continue
        kind, text = part
        r = p.add_run(text)
        r.bold = kind == 'b'
        r.italic = kind == 'i'
        if kind == 'c':
            r.font.name = 'Consolas'
            r.font.size = Pt(10.5)
    p.paragraph_format.space_after = Pt(6)
    return p


def code(text, key=None, caption=None, size=8.5):
    if caption:
        LST[key] = len(LST) + 1
        c = doc.add_paragraph()
        c.add_run(f'Listing {LST[key]}. ').bold = True
        c.add_run(caption).italic = True
        c.paragraph_format.space_after = Pt(2)
        c.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.rows[0].cells[0]
    shade(cell, 'F4F6FA')
    p = cell.paragraphs[0]
    lines = text.rstrip('\n').split('\n')
    for i, line in enumerate(lines):
        r = p.add_run(line)
        r.font.name = 'Consolas'
        r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Consolas')
        r.font.size = Pt(size)
        if i < len(lines) - 1:
            r.add_break()
    p.paragraph_format.space_after = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def figure(name, caption, width=16.5):
    FIG[name] = len(FIG) + 1
    doc.add_picture(str(IMG / name), width=Cm(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.add_run(f'Figure {FIG[name]}. ').bold = True
    c.add_run(caption).italic = True


def F(name):
    return f'figure {FIG[name]}'


def table(header, rows, key=None, caption=None, widths=None, font=10):
    if caption:
        TAB[key] = len(TAB) + 1
        c = doc.add_paragraph()
        c.add_run(f'Table {TAB[key]}. ').bold = True
        c.add_run(caption).italic = True
        c.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1, cols=len(header))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        r = cell.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(font)
        shade(cell, 'DCE3F0')
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            r = cells[i].paragraphs[0].add_run(str(val))
            r.font.size = Pt(font)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def src(rel):
    return (JAVA / rel).read_text(encoding='utf8')


def between(text, start, end=None, include_end=True):
    i = text.index(start)
    if end is None:
        return text[i:]
    j = text.index(end, i + len(start))
    return text[i:j + (len(end) if include_end else 0)]


def strip_package(text):
    return re.sub(r'^package .*?;\n\n(import .*?\n)+\n', '', text, flags=re.S)


def measurements(name):
    return json.loads((OUT / name).read_text(encoding='utf8').split('\n', 1)[1])


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


book = strip_package(src('domain/Book.java'))
author = strip_package(src('domain/Author.java'))
repo = strip_package(src('repository/BookRepository.java'))
svc = src('service/BookService.java')
ctrl = src('api/BookRestController.java')

# ---------------------------------------------------------------- 1. aim
h1('1. Aim of the work and repository')
para('To replace the in-memory list that has carried the catalogue since week 1 with a real database. '
     'By the end of the work the project stores books in PostgreSQL, the model has three entities joined '
     'by a one-to-many and a many-to-many association, the repository answers five derived queries and two '
     'written in JPQL, and the REST API of laboratory work no. 7 returns a page instead of a table.')
rich(('b', 'Repository, branch lab08: '), ('c', BRANCH))
rich(('b', 'Commit of the work: '), ('c', COMMIT),
     ' — "lab08: PostgreSQL, entities, associations, queries, paging", 95 files changed.')
rich(('b', 'Project: '), 'the work is done on top of the same project as labs 1–7 — ', ('i', 'Library Portal'),
     ' (Spring Boot 4.1.1, Java 25); branch lab08 was created from main. Nothing was rebuilt from scratch: the '
     'REST API, DTOs and ProblemDetail of lab 7, the aspects of lab 4 and the forms of lab 6 keep working; '
     'only the storage under BookService changed.')

h2('Environment and deviations from the handout')
table(['Component', 'Required', 'Used in this work'], [
    ['JDK', 'Java 25 LTS', 'java version "25.0.2" 2026-01-20 LTS'],
    ['Project from lab 7', 'spring-lab-01, branch main',
     'library-portal (the course project with labs 1–7), branch main → lab08'],
    ['PostgreSQL', '16 or newer, or Docker',
     'PostgreSQL 16.14 (official EDB binaries, local server on :5432); Docker is not installed'],
    ['Database and user', 'springlab / springlab', 'springlab / springlab; plus springlab_test for the tests'],
    ['Spring', '—', 'Spring Boot 4.1.1, Spring Data JPA 4.1.1, Hibernate ORM 7.4.5.Final, PostgreSQL JDBC 42.7.13'],
    ['Git', '2.40+', 'branch lab08 pushed to origin'],
], key='env', caption='Workstation', widths=[3.5, 4.5, 9])
para('Deviations: (1) the project is library-portal rather than spring-lab-01 — it is the same course project that '
     'already contains labs 1–7; (2) Book keeps the fields of labs 6–7 (available, publishedOn) next to the fields of '
     'the handout; (3) the API and the form still send the author as a name, so the service finds or creates the Author '
     'entity; (4) besides the console log, a small Hibernate StatementInspector (persistence/SqlRecorder) counts the SQL '
     'statements, so the numbers in the tables of tasks 3 and 4 are measured, not counted by eye.')

# ---------------------------------------------------------------- task 1
page_break()
h1('2. Task 1. PostgreSQL and the first entity')
h2('1.1. Branch and database')
figure('01_setup.png', 'Branch lab08 created from main; user and database springlab created in PostgreSQL 16.14')

h2('1.2. Dependencies')
pom = (ROOT / 'pom.xml').read_text(encoding='utf8')
code(between(pom, '        <!-- Lab 8: the data access layer', '<scope>runtime</scope>\n        </dependency>'),
     'pom', 'pom.xml — the starter and the driver; Spring Boot manages both versions')

h2('1.3. The datasource')
yml = (ROOT / 'src/main/resources/application.yml').read_text(encoding='utf8')
code(between(yml, '  # Lab 8: PostgreSQL', 'max-page-size: 50'), 'yml', 'application.yml — datasource and JPA settings')
para('open-in-view: false closes the persistence context when the service transaction ends, so nothing is loaded '
     'lazily while JSON or Thymeleaf is being rendered. ddl-auto: create-drop builds the schema from the entities at '
     'start-up and drops it on shutdown — for this class only; application-prod.yml sets ddl-auto: validate.')

h2('1.4. Book is an entity')
code(book, 'book', 'domain/Book.java — the entity (the record of labs 5–7 became a class)', size=8)
para('Why it is no longer a record: JPA needs a no-argument constructor, a non-final class (Hibernate creates proxy '
     'subclasses for lazy loading) and non-final fields that it can set after reading a row and later compare for dirty '
     'checking. A record is final, its fields are final and it has only the canonical constructor. The protected '
     'constructor exists for JPA only; our code uses Book(title, genre, year).')

h2('1.5. The repository replaces the in-memory list; the application starts against PostgreSQL')
figure('02a_startup_datasource.png', 'Start-up: Hibernate 7.4.5 connects to jdbc:postgresql://localhost:5432/springlab, '
                                     'PostgreSQL 16.14, PostgreSQLDialect — not H2, and no in-memory list is left')
figure('02b_startup_ddl.png', 'The DDL printed by show-sql: tables author, book, book_tag, tag and their foreign keys',
       width=14)
figure('04_psql_book.png', 'psql: \\dt and \\d book — the schema as the database reports it', width=15)
figure('06_psql_data.png', 'psql: select id, title, genre from book; author_id of every book; the rows of book_tag',
       width=15)
figure('08_api_page0.png', 'curl localhost:8080/api/books?page=0&size=3 still answers — now from the database',
       width=10)

# ---------------------------------------------------------------- task 2
page_break()
h1('3. Task 2. Associations')
figure('00_er_diagram.png', 'Entity diagram: Author 1—N Book, Book N—M Tag, with cardinalities and owning sides')
para('Written down before running anything: the author foreign key will be in table book (column author_id), because '
     'Book.author is @ManyToOne with @JoinColumn — the owning side; Author.books carries mappedBy and therefore no '
     'column; the many-to-many is stored in a third table book_tag(book_id, tag_id) declared by @JoinTable on Book.tags. '
     f'The schema in {F("02b_startup_ddl.png")}, {F("04_psql_book.png")} and the psql output below confirm it.')
code(between(book, '    // the owning side of book -> author', 'private Set<Tag> tags = new HashSet<>();'),
     'owning', 'Book — the owning side of both associations')
code(between(author, '    // ALL: saving or deleting', '    public Long getId()', include_end=False),
     'inverse', 'Author — the inverse side and the helper that keeps both sides in step')

h2('2.2. Seed data')
seeder = strip_package(src('persistence/DataSeeder.java'))
code(between(seeder, '    @Override', '        log.info'), 'seed',
     'persistence/DataSeeder.java — a CommandLineRunner for the development profiles', size=8)
para('Three authors, eight books, four tags. Joshua Bloch and Stephen Hawking own three books each; Effective Java, '
     'Java Concurrency in Practice, Spring in Action and A Brief History of Time carry two tags each. The runner is '
     'active in the profiles "default" (a run without a profile) and "dev"; prod never gets test data. The run issued '
     '27 INSERT statements: 4 tag + 3 author + 8 book (cascade = ALL from the author) + 12 book_tag (written at commit).')
figure('03_seed_sql.png', 'The INSERTs of the seed: every book is inserted with author_id already set', width=12)

h2('2.3. What the database shows')
figure('05_psql_assoc.png', 'psql: \\d book_tag, \\d author, \\d tag', width=14)
table(['Question', 'Your answer', 'How you checked it'], [
    ['Which tables were created?', 'author, book, book_tag, tag',
     f'\\dt ({F("04_psql_book.png")}) and the create table statements ({F("02b_startup_ddl.png")})'],
    ['Where is the author foreign key?', 'book.author_id bigint → author(id), constraint fkklnrv3weler2ftkweewlky958',
     '\\d book, section "Foreign-key constraints"'],
    ['What columns does book_tag have?',
     'book_id bigint not null, tag_id bigint not null; primary key (book_id, tag_id); both are foreign keys',
     f'\\d book_tag ({F("05_psql_assoc.png")})'],
    ['Does author have a books column?',
     'No — only id, country, name. Author.books is mappedBy = "author", the inverse side',
     f'\\d author ({F("05_psql_assoc.png")})'],
], key='t23', caption='Task 2.3 — filled in from what the database shows', widths=[4.2, 7.3, 5.5])
para('The two mistakes from the handout were checked against the SQL. Without mappedBy JPA would treat Author.books as '
     'a second owning side and create a third table author_books. Adding a book only with author.getBooks().add(book) '
     'would insert author_id = null, because JPA reads only the owning side. Author.addBook() sets both sides — which is '
     f'why the inserts in {F("03_seed_sql.png")} already carry author_id.')

# ---------------------------------------------------------------- task 3
page_break()
h1('4. Task 3. Queries')
code(repo, 'repo', 'repository/BookRepository.java — derived queries, JPQL, bulk update, N+1 cure, paging, Window',
     size=7.5)

h2('3.3. Every method called once')
para('All methods were called through the temporary endpoint POST /api/lab8/queries (api/Lab8Controller → '
     'persistence/QueryLab). Every call runs in its own transaction and returns the rows and the exact SQL that '
     'Hibernate sent.')
q = measurements('12_lab8_queries.json.txt')
table(['Method', 'Rows returned', 'Number of SQL statements'],
      [[m['method'], m['rows'], m['statements']] for m in q],
      key='t33', caption='Task 3.3 — measured with POST /api/lab8/queries', widths=[9, 3.5, 4.5])
figure('12_queries.png', 'POST /api/lab8/queries: results, statement counts and the SQL of every method')
top3 = next(m for m in q if m['method'].startswith('findTop3'))
sql = ' '.join(top3['sql'][0].split())
for kw in (' from ', ' where ', ' order by ', ' fetch '):
    sql = sql.replace(kw, '\n' + kw.strip() + ' ')
code(sql.replace(',', ', '), 'top3', 'SQL that Hibernate generated for findTop3ByGenreOrderByYearDesc(PROGRAMMING)')
para('Subject "findTop3" → fetch first 3 rows only; predicate "ByGenre" → where genre = ?; "OrderByYearDesc" → '
     'order by year desc. The JPQL that Spring Data builds is select b from Book b where b.genre = ?1 order by '
     'b.year desc, executed with maxResults = 3.')
para('The bulk update archiveOlderThan(2000) is a single statement — update book set archived=true where year<? — and '
     'changed 1 row (A Brief History of Time, 1988). It needs @Modifying and a read-write transaction '
     '(BookService.archiveOlderThan is @Transactional).')

h2('3.4. The N+1 exercise')
n1 = measurements('13_lab8_n_plus_one.json.txt')
table(['Variant', 'Books', 'SQL statements'], [[m['method'], m['rows'], m['statements']] for m in n1],
      key='t34', caption='Task 3.4 — N+1 before and after', widths=[9, 3, 4])
figure('13_n_plus_one.png', 'GET /api/lab8/n-plus-one: 4 statements before the fetch join, 1 after (and 1 with '
                            '@EntityGraph)')
rich(('b', 'Before the fetch join: 4 statements. After: 1. '),
     'findAll() loads the 8 books with author_id only (author is LAZY). Each call book.getAuthor().getName() '
     'initialises a proxy with select ... from author where id=?. There are 3 distinct authors, and the persistence '
     'context returns an already loaded author from its first-level cache, so the count is 1 + 3 = 4, not 1 + 8; with '
     '1 000 books by 300 authors it would be 301. select b from Book b join fetch b.author brings the author columns in '
     'the same row, so it is 1 statement; @EntityGraph(attributePaths = "author") on findByGenre does the same '
     'declaratively (with a left join).')

# ---------------------------------------------------------------- task 4
page_break()
h1('5. Task 4. Pagination and sorting')
code(between(ctrl, '    @GetMapping\n    public Page<BookResponse> list(', '.map(mapper::toResponse);\n    }'),
     'list', 'BookRestController — the collection endpoint is paged (@PageableDefault size 5, sort by title)')
code(between(svc, '    @Audited(action = "BOOK_LIST", logArguments = true)', '        return page;\n    }'),
     'find', 'BookService.find — the transaction boundary; the page brings its authors, the tags come in one batch')
para('PagingConfig enables @EnableSpringDataWebSupport(pageSerializationMode = VIA_DTO), so the JSON shape is stable: '
     'content plus page {size, number, totalElements, totalPages}. spring.data.web.pageable.max-page-size = 50 '
     'prevents ?size=100000 from asking for the whole table.')

h2('4.2. The checks')
figure('09a_api_page1_json.png', 'GET /api/books?page=1&size=3&sort=year,desc — the second page, newest first, '
                                 'totalElements = 8, totalPages = 3', width=10.5)
txt = (OUT / '09_api_page1_sort.txt').read_text(encoding='utf8')
code(txt.split('----- SQL')[0].strip(), 'json', 'The JSON of the paged response', size=7.5)
figure('09b_api_page1_sql.png', 'The SQL of that request: the page query (offset / fetch first), the count query, and '
                                'the batched select of the tags', width=10)
para('The paged request caused (1) select ... from book left join author ... order by year desc offset ? rows fetch '
     'first ? rows only — the page itself; (2) select count(b1_0.id) from book — the count query. It is there because a '
     'Page promises totalElements and totalPages, which the page query alone cannot know. The third statement '
     '(select ... from book_tag join tag where book_id = any(?)) is not part of paging: it loads the tags of all three '
     'books at once (@BatchSize) instead of one select per book.')

h2('sort=nosuchfield')
figure('10a_nosuchfield_500.png', 'Before: 500 Internal Server Error, PropertyReferenceException "No property '
                                  'nosuchfield found for type Book"')
figure('10b_nosuchfield_400.png', 'After the handler in ApiExceptionHandler: 400, application/problem+json, type '
                                  'unknown-sort-property')
para('The sort parameter is only a string until Spring Data resolves it against the entity while it builds the query — '
     'at run time, inside SimpleJpaRepository (QueryUtils.toOrders → PropertyPath.from). The compiler never sees it, so '
     'it can only fail at run time. The project now maps PropertyReferenceException to a 400 problem detail: it is the '
     'client who asked for a property that does not exist.')
code(between(src('api/ApiExceptionHandler.java'), '    @ExceptionHandler(PropertyReferenceException.class)',
             '        return problem;\n    }'), 'nosuch', 'ApiExceptionHandler — unknown sort property → 400')

h2('4.3. Page against Slice')
ps = measurements('14_lab8_page_vs_slice.json.txt')
table(['Return type', 'Rows', 'SQL statements', 'What the client learns'], [
    ['Page<Book> findAll(pageable)', ps[0]['rows'], ps[0]['statements'], 'content + totalElements = 8, totalPages = 3'],
    ['Slice<Book> findSliceBy(pageable)', ps[1]['rows'], ps[1]['statements'], 'content + hasNext = true'],
], key='t43', caption='Task 4.3 — the same request (page 1, size 3, year desc)', widths=[5.5, 1.5, 3, 7])
figure('14_page_vs_slice.png', 'GET /api/lab8/paging: a Page costs select + count, a Slice only the select')
para('The Slice asks for size + 1 rows (fetch first 4) and uses the extra row only to know whether there is a next '
     'page. Instead of changing the return type back and forth, the comparison uses a second repository method, '
     'findSliceBy, so both numbers come from the same run. What I would publish: a Page for this catalogue API — '
     'clients show "page 2 of 3", and the count over a small indexed table is cheap. For an infinite feed or a very '
     'large table I would publish a Slice (or keyset scrolling, see variant 7), because the count query can cost as much '
     'as reading the whole table.')

h2('4.4. Commit and push')
figure('21_git.png', 'git commit and git push -u origin lab08', width=12)
figure('25_github_lab08.png', 'The commit on GitHub, branch lab08')

# ---------------------------------------------------------------- what else the ORM does
page_break()
h1('6. What else the ORM did (used in the answers)')
figure('17_post.png', 'POST /api/books: existing tags found, a new author inserted, the book inserted with author_id, '
                      'one book_tag row per tag')
figure('18_put_dirty.png', 'PUT /api/books/1: no save() in the service — dirty checking issues the UPDATE '
                           '"where id=? and version=?"; version became 1')
code(between(svc, '    @Transactional\n    @Audited(action = "BOOK_UPDATE"', '        return book;\n    }'),
     'update', 'BookService.update — changes a managed entity and never calls save()')
figure('19_delete_orphan.png', 'DELETE /api/books/9: author.removeBook(book) → orphanRemoval deletes the book_tag '
                               'rows and the book "where id=? and version=?"')
figure('11_lazy_init.png', 'LazyInitializationException met during the work: with open-in-view off, the controller '
                           'touched a lazy author after the transaction had ended')
figure('21_catalogue_page.png', 'The Thymeleaf catalogue is paged too: /books?size=3&page=1 — "Page 2 of 3"', width=13)
figure('23_swagger.png', 'Swagger UI: the paged GET /api/books, GET /api/books/scroll and the Lab 8 group', width=13)

# ---------------------------------------------------------------- variant 7
page_break()
h1('7. Individual variant 7 — Window and keyset scrolling')
para('Requirement: Window<Book> with a ScrollPosition; compare the SQL with offset paging.')
code(between(repo, '    // ---------------- Individual variant 7',
             '    Window<Book> findBy(ScrollPosition position, Sort sort, Limit limit);'),
     'window', 'BookRepository — the Window method')
code(between(svc, '    /**\n     * Lab 8, variant 7: the first window',
             'ScrollPosition.forward(Map.of("year", Integer.MAX_VALUE, "id", Long.MAX_VALUE));')
     + '\n\n' + between(svc, '    public Window<Book> scroll(', '        return window;\n    }'),
     'scroll', 'BookService.scroll — newest first; the id is the tie-breaker of the sort keys')
code(between(ctrl, '    @GetMapping("/scroll")',
             '        return new BookWindowResponse(content, window.hasNext(), next);\n    }'),
     'scrollapi', 'BookRestController.scroll — GET /api/books/scroll?size=3&afterYear=&afterId=')
figure('16_scroll_api.png', 'Three windows of GET /api/books/scroll: each "next" cursor is the (year, id) of the last '
                            'book; the third window has hasNext = false', width=15)
sc = measurements('15_lab8_scroll.json.txt')
table(['Request', 'Books', 'SQL', 'Key part of the SQL'], [
    ['Keyset, first window', sc[0]['rows'], sc[0]['statements'],
     'where year<? or year=? and id<? order by year desc, id desc fetch first ? rows only'],
    ['Keyset, after (2018, id 1)', sc[1]['rows'], sc[1]['statements'],
     'where year<? or year=? and id<? ... fetch first ? rows only'],
    ['Offset Page, page 1, size 3', sc[2]['rows'], sc[2]['statements'],
     'order by year desc, id desc offset ? rows fetch first ? rows only + select count(...)'],
], key='t7', caption='Variant 7 — keyset against offset (GET /api/lab8/scroll)', widths=[4, 1.4, 1.2, 10.4])
figure('15_scroll_compare.png', 'GET /api/lab8/scroll: the same second "page" by keyset and by offset')
para('Comparison. Offset paging makes the database produce the first page × size rows, throw them away and return the '
     'next size rows: the deeper the page, the more rows are read, and a row inserted or deleted between two requests '
     'shifts the pages (a book is shown twice or skipped). The Page version also adds a count query. Keyset scrolling '
     'remembers the sort keys of the last row and asks only for the rows "after" them: with an index on (year, id) the '
     'database seeks straight to the position, the cost does not grow with depth, concurrent inserts do not shift the '
     'window, and there is no count. The price: no "jump to page 7" and no total — only "next".')
rich(('b', 'A pitfall found while doing the variant. '),
     'The first version passed ScrollPosition.keyset() (the empty initial position) for the first window. The second '
     'window then returned the same three books, and its SQL had no where clause at all. After a restart, asking for a '
     'window with a cursor first produced the correct "where year<? or year=? and id<?" — so Spring Data JPA 4.1.1 '
     'caches the query it renders for the first call of the method, and the query rendered for the empty position '
     '(without the keyset predicate) was reused for every later position. The fix: never pass the empty position — the '
     'first window starts "after" a sentinel (Integer.MAX_VALUE, Long.MAX_VALUE), so every call has the same query '
     'shape. The test Lab8RepositoryTests.keysetScrolling walks the whole catalogue window by window and asserts that no '
     'book repeats and that no statement contains OFFSET.')

# ---------------------------------------------------------------- tests
h1('8. Automated checks')
figure('20_tests.png', './mvnw test — 41 tests, 0 failures (9 new ones in Lab8RepositoryTests), run against the '
                       'PostgreSQL database springlab_test', width=15)
table(['Test (Lab8RepositoryTests)', 'What it proves'], [
    ['derivedQueries', 'the five derived queries return rows from the seed'],
    ['percentIsEscaped', 'findByTitleContainingIgnoreCase("%") matches nothing — the % is escaped'],
    ['jpqlAndBulkUpdate', 'JPQL search with named parameters; the bulk update archives 1 row; '
                          'clearAutomatically makes the next read see it'],
    ['owningSideWritesTheForeignKey', 'after create(...) the row in book has the author_id of Joshua Bloch '
                                      'and two book_tag rows'],
    ['nPlusOne', 'findAll + getAuthor = 1 + number of authors statements; findAllWithAuthor = 1'],
    ['pageAndSlice', 'Page = 2 statements (one is select count), Slice = 1'],
    ['keysetScrolling', 'windows cover the catalogue without repeats, no OFFSET in the SQL'],
    ['pagedEndpoint', 'GET /api/books?page=1&size=3&sort=year,desc → 3 items, number 1, the total'],
    ['unknownSortPropertyIs400', 'sort=nosuchfield → 400 problem detail'],
], widths=[5.5, 11.5], font=9.5)

# ---------------------------------------------------------------- checklist
page_break()
h1('9. Demonstration checklist (section 7 of the handout)')
table(['No.', 'What is shown', 'Where (report / application)'], [
    [1, 'The application starting against PostgreSQL',
     f'{F("02a_startup_datasource.png")}: jdbc:postgresql, PostgreSQL 16.14; BookRepository extends JpaRepository'],
    [2, 'The Book entity', f'Listing {LST["book"]}; @Id, @GeneratedValue(IDENTITY), protected constructor; answers 1–2'],
    [3, 'The tables in the database', f'{F("02b_startup_ddl.png")}, {F("04_psql_book.png")}, {F("05_psql_assoc.png")}'],
    [4, 'The two associations and their owning sides',
     f'Listings {LST["owning"]}–{LST["inverse"]}, {F("00_er_diagram.png")}; answer 5'],
    [5, 'The author_id column and the book_tag table', f'Table {TAB["t23"]}, {F("05_psql_assoc.png")}'],
    [6, 'The addBook helper', f'Listing {LST["inverse"]}; answer 8'],
    [7, 'Three derived queries running, SQL read aloud',
     f'POST /api/lab8/queries — table {TAB["t33"]}, {F("12_queries.png")}, listing {LST["top3"]}'],
    [8, 'The JPQL query and the bulk update', 'search(:genre, :year), archiveOlderThan with @Modifying; answer 15'],
    [9, 'The N+1 count, before and after', f'Table {TAB["t34"]}, {F("13_n_plus_one.png")}: 4 → 1'],
    [10, 'GET /api/books?page=1&size=3&sort=year,desc', f'{F("09a_api_page1_json.png")}: 3 books, year desc, total 8'],
    [11, 'The SQL of the paged query', f'{F("09b_api_page1_sql.png")}: the page query and select count'],
    [12, 'The individual variant in action', f'Section 7, {F("16_scroll_api.png")}, table {TAB["t7"]}'],
], widths=[1, 6, 10], font=9.5)
para('How to run the demonstration: start PostgreSQL (pg_ctl -D C:\\Users\\kuatw\\tools\\pgdata start), run '
     './mvnw spring-boot:run, open psql -U springlab -d springlab (\\d book, \\d book_tag), and call '
     '/api/books?page=1&size=3&sort=year,desc, POST /api/lab8/queries, /api/lab8/n-plus-one, /api/lab8/paging, '
     '/api/lab8/scroll and /api/books/scroll; the console shows the SQL (show-sql + format_sql).')

# ---------------------------------------------------------------- questions
page_break()
h1('10. Answers to the questions of the demonstration (section 10 of the handout)')
QA = [
    ('Why can your Book no longer be a record? What exactly does JPA require of the class?',
     'JPA requires: @Entity; a primary key marked with @Id; a public or protected no-argument constructor; a top-level '
     'class that is not final, with persistent fields that are not final. Hibernate instantiates the entity with the '
     'no-arg constructor and then sets the fields from the row; it creates proxy subclasses for lazy references; it '
     'changes the fields of managed entities and compares them at flush. A record is implicitly final, its fields are '
     'final and it has only the canonical constructor, so none of this is possible and the application would fail at '
     'start-up. Book is now a class with a protected constructor for JPA and a public Book(title, genre, year) for '
     'the code.'),
    ('What does @GeneratedValue(strategy = IDENTITY) do, and where does the value come from?',
     'The value is generated by the database: the column is created as "id bigint generated by default as identity" '
     f'({F("04_psql_book.png")}) and PostgreSQL takes the next value of the identity sequence on INSERT. Hibernate does '
     'not know the id before the insert, so persist()/save() runs the INSERT immediately and reads the generated key '
     'back (that is why IDENTITY inserts cannot be batched). Until then the id is null.'),
    ('Why is your enum mapped with EnumType.STRING? What would break with ORDINAL?',
     'STRING stores the name: genre varchar(20) with values like PROGRAMMING, and Hibernate even adds a check constraint '
     'genre in (NOVEL, SCIENCE, ...). ORDINAL stores the position (0, 1, 2 ...). If a constant is inserted in the middle '
     'or the enum is reordered, every existing row silently changes its meaning (2 was PROGRAMMING, now it is HISTORY), '
     'and the data cannot be read without the Java source. With STRING only a rename is dangerous, and it fails loudly.'),
    ('What is the @Version field for? What happens when two transactions change the same row?',
     'It is optimistic locking. Every UPDATE and DELETE is written as "... where id=? and version=?" and increments the '
     f'version ({F("18_put_dirty.png")}: version 0 → 1). If two transactions both read version 0 and both update, the '
     'first commits and makes it 1; the second one\'s "where version=0" matches 0 rows, Hibernate throws '
     'OptimisticLockException (in Spring: ObjectOptimisticLockingFailureException) and that transaction is rolled back. '
     'The lost update is detected instead of silently overwriting the first change; the client reloads and retries '
     '(an API would answer 409 Conflict).'),
    ('Which side of the book–author association owns it? How do you know from the code alone?',
     'Book.author owns it: it is the @ManyToOne side with @JoinColumn(name = "author_id"), the side that maps the column. '
     'Author.books has mappedBy = "author", which literally says "this association is mapped by the field author of '
     'Book": an inverse side has mappedBy and no column. In a @ManyToOne/@OneToMany pair the @ManyToOne side always owns, '
     'because the foreign key lives in the table of the "many".'),
    ('What does mappedBy mean? What would happen if you removed it from Author.books?',
     'mappedBy names the field of the other entity that owns the association, so this collection is only a mirror and '
     'JPA ignores it when writing. Without mappedBy JPA would see two independent associations: Book.author with '
     'author_id, and a unidirectional @OneToMany Author.books, which by default is mapped with a join table — Hibernate '
     'would create a third table author_books(author_id, books_id) and store every link twice, in two places that can '
     'disagree.'),
    ('Which table holds the many-to-many? What are its columns, and who created it?',
     f'book_tag ({F("05_psql_assoc.png")}): book_id bigint not null → book(id), tag_id bigint not null → tag(id), '
     'primary key (book_id, tag_id). Hibernate created it (ddl-auto: create-drop) from @JoinTable(name = "book_tag", '
     'joinColumns = book_id, inverseJoinColumns = tag_id) on Book.tags; there is no entity class for it.'),
    ('Why does addBook set both sides? Show what the SQL does if it sets only one.',
     'Because the Java objects and the database are two different worlds. The database needs book.author_id, and JPA '
     'reads it only from the owning side, Book.author. The objects need author.books to contain the book, or the same '
     'transaction sees an inconsistent graph. If only author.getBooks().add(book) is called, the insert is '
     '"insert into book (..., author_id, ...) values (..., null, ...)" and the link is lost. If only book.setAuthor(author) '
     'is called, the database is right but author.getBooks() in the same persistence context does not contain the book. '
     'addBook does both; in the project setAuthor is package-private, so the rest of the code cannot skip the helper.'),
    ('What does CascadeType.ALL on Author.books cause when an author is deleted?',
     'Every operation on the author is repeated on its books: persist (saving an author saves its new books — the seed '
     'saves only authors), merge, refresh, detach and remove. Deleting an author therefore first deletes all its books '
     '(and their book_tag rows) and then the author; without the cascade the delete would fail on the foreign key '
     'book.author_id.'),
    ('What is orphanRemoval = true doing that the cascade does not?',
     'A cascade reacts to an operation on the parent. orphanRemoval reacts to a child being taken out of the parent\'s '
     'collection while the parent stays: a book removed from author.books is an orphan and is deleted. In the project '
     'DELETE /api/books/9 calls author.removeBook(book) and never calls repository.delete; '
     f'{F("19_delete_orphan.png")} shows the resulting "delete from book_tag where book_id=?" and '
     '"delete from book where id=? and version=?".'),
    ('Where is the implementation of BookRepository? Who wrote it and when?',
     'Nobody wrote it by hand. At start-up Spring Data JPA (JpaRepositoryFactory) creates a JDK dynamic proxy for the '
     f'interface (jdk.proxy4.$Proxy188 in the stack trace of {F("10a_nosuchfield_500.png")}). Calls to the CRUD and '
     'paging methods go to the generic SimpleJpaRepository; for every other method Spring Data parses the name into a '
     'PartTree once, builds the query and executes it through the EntityManager; @Query methods use the given JPQL.'),
    ('Split findTop3ByGenreOrderByYearDesc into subject and predicate and read out the JPQL.',
     'Subject: "find" (a query) + "Top3" (limit 3). "By" separates the subject from the predicate. Predicate: "Genre" '
     '(genre = ?1). Order clause: "OrderBy" "Year" "Desc". JPQL: select b from Book b where b.genre = ?1 order by '
     'b.year desc, with maxResults 3. SQL: select ... from book b1_0 where b1_0.genre=? order by b1_0.year desc fetch '
     f'first ? rows only (listing {LST["top3"]}).'),
    ('What SQL does findByTitleContainingIgnoreCase produce? What happens to a % in the argument?',
     'select ... from book b1_0 where upper(b1_0.title) like upper(?) escape \'\\\'. Spring Data wraps the argument as '
     '%argument% and escapes the special characters % and _ inside it with the escape character \\, so a % in the '
     'argument is a literal percent sign, not a wildcard: findByTitleContainingIgnoreCase("%") returned 0 rows '
     f'(table {TAB["t33"]}), while an unescaped "%%%" would have matched all 8 books.'),
    ('Why did you use named parameters in the JPQL query rather than ?1?',
     ':genre and :year are bound by name with @Param, so the query reads by itself, the order of the method parameters '
     'can change without breaking it, the same parameter can be used twice, and a mismatch is reported at start-up. '
     '?1/?2 depend on positions and break silently when someone reorders the arguments. Both are bind parameters — '
     'neither is string concatenation — so both are safe from SQL injection.'),
    ('Why does the bulk update need @Modifying? What does clearAutomatically = true prevent?',
     'Spring Data runs an @Query with getResultList()/getSingleResult() by default; an UPDATE must be run with '
     'executeUpdate(), and @Modifying says so (without it: "Not supported for DML operations"). It also needs a '
     'transaction (else TransactionRequiredException) — BookService.archiveOlderThan is @Transactional. A bulk update '
     'goes straight to the database and bypasses the persistence context, so entities already loaded keep the old '
     'archived = false. clearAutomatically = true clears the context after the update, so the next read loads the fresh '
     'row; the test jpqlAndBulkUpdate checks exactly that.'),
    ('Show the N+1 numbers. Explain where the extra queries came from and why the fetch join removes them.',
     f'Before: 4 statements, after: 1 (table {TAB["t34"]}, {F("13_n_plus_one.png")}). findAll() selected the 8 books '
     'with author_id only, because author is LAZY. Touching book.getAuthor().getName() initialised the proxy with one '
     'select per author; there are 3 authors and the first-level cache serves the repeats, hence 1 + 3. "join fetch '
     'b.author" makes Hibernate select the author columns in the same statement and fill the association from the join, '
     'so no proxy is left to initialise.'),
    ('What is the difference between join and join fetch in JPQL?',
     'join only joins for filtering or ordering (where a.name = ...): the selected Book still has an uninitialised lazy '
     'author, and touching it costs a select. join fetch also selects the columns of the joined entity and initialises '
     'the association in the result, so the data comes in one statement. Fetching a collection together with paging '
     'makes Hibernate page in memory, which is why the page query fetches only the author and loads the tags with '
     '@BatchSize.'),
    ('How many SQL statements does a Page cause? And a Slice? Why the difference?',
     f'Page: 2 — the select with offset/fetch first and select count(...) (table {TAB["t43"]}). Slice: 1 — it fetches '
     'size + 1 rows and uses the extra row only to set hasNext. A Page must report totalElements and totalPages, which '
     'only a count can give; a Slice promises only "is there a next one". (Spring Data skips the count when the first '
     'page is not full.)'),
    ('What happened when you passed sort=nosuchfield? Why is that a run-time error?',
     'At first: 500 Internal Server Error with PropertyReferenceException "No property nosuchfield found for type Book" '
     f'({F("10a_nosuchfield_500.png")}). The sort parameter is a string in the URL; Spring MVC puts it into the Pageable '
     'without knowing the entity, and only Spring Data resolves it against Book while building the query, inside the '
     'repository call. Nothing is compiled from it, so it cannot fail earlier. The project now answers 400 with a '
     f'problem detail ({F("10b_nosuchfield_400.png")}).'),
    ('What is ddl-auto set to in your project? What is the default on PostgreSQL, and what should it be in production?',
     'create-drop in application.yml (this class, the default and dev runs), create in the tests (several test contexts '
     'share springlab_test), validate in application-prod.yml. Spring Boot\'s default is create-drop only for an '
     'embedded database (H2, HSQLDB, Derby); for PostgreSQL it is none — the "relation book does not exist" trap. In '
     'production: none or validate, with the schema created and changed by versioned migrations (Flyway or Liquibase); '
     'Hibernate must never drop or alter production tables.'),
    ('Why did you set open-in-view: false? What would you have to do differently with it on?',
     'With open-in-view on (Spring Boot\'s default, which logs a warning) the EntityManager stays open for the whole '
     'HTTP request, so lazy associations load silently during JSON serialisation or Thymeleaf rendering — hidden N+1 '
     'queries outside the service, and a database connection held while the view renders. With it off, the service '
     'transaction is the boundary: whatever the response needs must be fetched inside it (entity graph, join fetch, '
     f'@BatchSize + initialisation in BookService.find); otherwise — LazyInitializationException ({F("11_lazy_init.png")}). '
     'With it on those fetches would not be needed for correctness, but I would have to watch the SQL log for the '
     'hidden queries.'),
]
for i, (question, answer) in enumerate(QA, 1):
    p = doc.add_paragraph()
    p.add_run(f'{i}. {question}').bold = True
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(2)
    doc.add_paragraph(answer).paragraph_format.space_after = Pt(8)

# ---------------------------------------------------------------- conclusions
h1('11. Conclusions')
para('The ORM took over everything between the objects and the tables: Hibernate created the four tables, the foreign '
     'keys and the join table from annotations, wrote the INSERTs and UPDATEs itself (dirty checking, the version check, '
     'orphan removal), and Spring Data turned method names into SQL — no SQL was written by hand except two JPQL queries. '
     'The REST API of lab 7 became paged and sorted with one Pageable parameter, and the individual variant added keyset '
     'scrolling, which reads only the rows it returns and needs no count.')
para('What it took away is direct control and visibility: the number of statements depends on fetch types and on what '
     'the code touches (4 statements instead of 1 for N+1), a Page silently adds a count query, a missing transaction '
     'with open-in-view off turns into LazyInitializationException, a wrong sort property is found only at run time, and '
     'even the framework can surprise (the cached keyset query). Reading — and measuring — the SQL that Hibernate prints '
     'is therefore part of using an ORM, not an optional extra.')

target = LAB / 'RWPSF_Lab8_IT1-2409IS_Bektursyn.docx'
doc.save(target)
print(target, 'figures:', len(FIG), 'tables:', len(TAB), 'listings:', len(LST))
