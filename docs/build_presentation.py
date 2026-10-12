"""Generates docs/Defense_Presentation.pptx (English, python-pptx). Run: python docs/build_presentation.py"""
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# ---- palette: Spring green on deep forest ----
DARK = RGBColor(0x14, 0x2B, 0x22)      # deep forest (title / closing slides, code cards)
GREEN = RGBColor(0x6D, 0xB3, 0x3F)     # Spring green — the accent
SOFT = RGBColor(0xEA, 0xF4, 0xE2)      # light green tint for cards
INK = RGBColor(0x1D, 0x26, 0x21)       # body text
MUTED = RGBColor(0x5E, 0x6B, 0x63)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xC9, 0xD8, 0xC0)
CODE_TXT = RGBColor(0xD8, 0xF0, 0xC8)
RED = RGBColor(0xC0, 0x39, 0x2B)

HEAD = "Cambria"
BODY = "Calibri"
MONO = "Courier New"

W, H = 13.333, 7.5
REPO = "github.com/AkAyu44/library-portal"

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]
slide_no = 0


def text(slide, x, y, w, h, runs, size=16, color=INK, bold=False, font=BODY,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.0, space_after=0, line=None):
    """runs: str, or list of paragraphs; a paragraph is str or list of (text, {opts})."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(margin))
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        if line:
            p.line_spacing = line
        parts = para if isinstance(para, list) else [(para, {})]
        for t, o in parts:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", False)
            f.color.rgb = o.get("color", color)
    return box


def bullets(slide, x, y, w, h, items, size=16, color=INK, space=8):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(0))
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space)
        parts = item if isinstance(item, list) else [(item, {})]
        # a real PowerPoint bullet (green •) with a hanging indent
        ppr = p._p.get_or_add_pPr()
        ppr.set("marL", str(Inches(0.3)))
        ppr.set("indent", str(-Inches(0.3)))
        ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
        clr = ppr.makeelement(ns + "buClr", {})
        clr.append(clr.makeelement(ns + "srgbClr", {"val": str(GREEN)}))
        ppr.append(clr)
        ppr.append(ppr.makeelement(ns + "buFont", {"typeface": "Arial"}))
        ppr.append(ppr.makeelement(ns + "buChar", {"char": "•"}))
        for t, o in parts:
            r = p.add_run()
            r.text = t
            r.font.name = o.get("font", BODY)
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", False)
            r.font.color.rgb = o.get("color", color)
    return box


def rect(slide, x, y, w, h, fill, rounded=True, line=None, shadow=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        shp.adjustments[0] = 0.08
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    if not shadow:
        shp.shadow.inherit = False
    return shp


def circle(slide, x, y, d, fill, label, size=16, color=WHITE):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(0))
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.name = HEAD
    r.font.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = color
    return shp


def arrow(slide, x1, y1, x2, y2, color=GREEN):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(2)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd",
                          {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c


def code(slide, x, y, w, h, lines, size=12, title=None):
    rect(slide, x, y, w, h, DARK)
    top = y + 0.2
    if title:
        text(slide, x + 0.3, y + 0.15, w - 0.6, 0.35, title, size=11, color=GREEN, bold=True)
        top = y + 0.55
    text(slide, x + 0.3, top, w - 0.6, h - (top - y) - 0.15, lines, size=size, color=CODE_TXT, font=MONO,
         line=1.05)


def new_slide(title=None, lab=None, variant=None, notes="", dark=False):
    global slide_no
    slide_no += 1
    s = prs.slides.add_slide(BLANK)
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = DARK if dark else WHITE
    if title:
        text(s, 0.6, 0.42, 9.6 if lab else 12.1, 0.8, title, size=34, bold=True, font=HEAD, color=DARK,
             anchor=MSO_ANCHOR.MIDDLE)
    if lab:
        pill = rect(s, 10.55, 0.55, 2.2, 0.55, GREEN)
        pill.adjustments[0] = 0.5
        label = f"Lab {lab}" + (f"  ·  V{variant}" if variant else "")
        text(s, 10.55, 0.55, 2.2, 0.55, label, size=15, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if not dark:
        text(s, 0.6, 7.0, 6.0, 0.3, "Library Portal  ·  RWPSF 3305  ·  Kuat Bektursyn", size=10, color=MUTED)
        text(s, 11.73, 7.0, 1.0, 0.3, str(slide_no), size=10, color=MUTED, align=PP_ALIGN.RIGHT)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def table(slide, x, y, w, rows, col_w, size=14, row_h=0.42, header_fill=DARK):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w),
                                   Inches(row_h * len(rows)))
    tbl = shape.table
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        tbl.rows[r].height = Inches(row_h)
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if r == 0 else (SOFT if r % 2 == 0 else WHITE)
            cell.margin_left = cell.margin_right = Inches(0.12)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.paragraphs[0].text = ""
            run = tf.paragraphs[0].add_run()
            run.text = str(val)
            run.font.name = BODY
            run.font.size = Pt(size)
            run.font.bold = r == 0 or c == 0
            run.font.color.rgb = WHITE if r == 0 else INK
    return tbl



# =====================================================================
# 1. Title
s = new_slide(dark=True, notes=(
    "Good afternoon. I am presenting the final educational project for the course Web Application Development "
    "with Spring Framework. It is one application, a library catalogue, which combines the topics and the "
    "individual variants of all six laboratory works."))
circle(s, 0.9, 1.3, 0.9, GREEN, "6", size=30)
text(s, 0.9, 2.45, 11.5, 1.2, "Library Portal", size=60, bold=True, font=HEAD, color=WHITE)
text(s, 0.9, 3.65, 11.5, 0.8, "Final educational project: six laboratory works in one Spring Boot application",
     size=24, color=CODE_TXT)
text(s, 0.9, 5.2, 8.0, 1.4, [
    [("Kuat Bektursyn", {"bold": True, "color": WHITE})],
    "Group IT1-2409IS  ·  RWPSF 3305 — Web Application Development with Spring Framework",
    "IITU, Almaty, 2026",
], size=16, color=CODE_TXT, space_after=4)
text(s, 8.9, 6.25, 3.8, 0.4, REPO, size=14, color=GREEN, align=PP_ALIGN.RIGHT)

# =====================================================================
# 2. Idea
s = new_slide("The idea: six labs as six layers", notes=(
    "All the labs were done in one repository, spring-lab-01, but each one on its own branch and as a separate "
    "demonstration. In the final project they become layers of one library catalogue: REST and MVC are the "
    "entry point, IoC assembles the application, configuration supplies the parameters, AOP adds cross-cutting "
    "behaviour and forms are the user interface. Everything works on one BookService, and every experiment from "
    "the reports is reproduced by an automated test."))
stats = [("6", "laboratory works\nin one application"), ("17", "integration tests —\none per experiment"),
         ("6", "individual variants\nimplemented"), ("1", "shared BookService\nfor every layer")]
for i, (num, lbl) in enumerate(stats):
    x = 0.6 + i * 3.08
    rect(s, x, 1.65, 2.83, 2.25, SOFT)
    text(s, x, 1.8, 2.83, 1.1, num, size=60, bold=True, font=HEAD, color=GREEN, align=PP_ALIGN.CENTER)
    text(s, x + 0.15, 2.95, 2.53, 0.85, lbl, size=14, color=INK, align=PP_ALIGN.CENTER)
text(s, 0.6, 4.3, 5.0, 0.4, "Technology stack", size=20, bold=True, font=HEAD, color=DARK)
bullets(s, 0.6, 4.8, 5.8, 2.0, [
    "Java 25, Spring Boot 4.1.1 (Spring Framework 7), Maven 3.9",
    "Spring MVC + Thymeleaf, Hibernate Validator",
    "Spring AOP (AspectJ annotations), Spring Boot Actuator",
], size=15)
text(s, 6.9, 4.3, 5.8, 0.4, "Domain", size=20, bold=True, font=HEAD, color=DARK)
bullets(s, 6.9, 4.8, 5.8, 2.0, [
    "Book catalogue: list, details, add, delete",
    "REST API /api/books and HTML pages /books",
    "Two interface languages — EN and RU",
], size=15)

# =====================================================================
# 3. Architecture
s = new_slide("Architecture: where each lab lives", notes=(
    "A request enters the DispatcherServlet. The LocaleChangeInterceptor from lab 6 selects the language. "
    "The REST controllers from labs 1, 3 and 5 return JSON, the page controller returns a Thymeleaf view name. "
    "Errors are split between two ControllerAdvice classes. Below is the CGLIB proxy of BookService, through "
    "which the four aspects from lab 4 pass. The service itself uses the repository, the notifications from "
    "lab 2 and the citation cache, whose parameters come from the configuration of lab 3."))


def node(x, y, w, h, title, sub, lab, fill=SOFT, tc=DARK):
    rect(s, x, y, w, h, fill)
    text(s, x + 0.2, y + 0.12, w - 0.4, 0.4, title, size=15, bold=True, color=tc)
    text(s, x + 0.2, y + 0.5, w - 0.4, h - 0.55, sub, size=12, color=tc if fill == DARK else MUTED)
    circle(s, x + w - 0.5, y - 0.18, 0.42, GREEN, lab, size=11)


node(0.6, 1.65, 3.6, 1.1, "DispatcherServlet", "LocaleChangeInterceptor · ?lang=ru", "6")
node(4.8, 1.65, 3.8, 1.1, "@RestController", "Book, Hello, Config, System API", "1·5")
node(9.2, 1.65, 3.5, 1.1, "@Controller + Thymeleaf", "BookPageController · /books", "5·6")
node(4.8, 3.35, 7.9, 1.1, "CGLIB proxy of BookService", "Audit(1) → Logging(2) → Timing(3) → CallTracing",
     "4", fill=DARK, tc=WHITE)
node(0.6, 3.35, 3.6, 1.1, "@ControllerAdvice", "HTML 404  ·  ProblemDetail", "6")
node(0.6, 5.15, 3.6, 1.1, "BookRepository", "in-memory, ConcurrentHashMap", "5")
node(4.8, 5.15, 3.8, 1.1, "NotificationService", "@Primary EmailNotifier ← app.mail", "2·3")
node(9.2, 5.15, 3.5, 1.1, "CitationCache", "app.cache.ttl / max-entries", "3")
arrow(s, 4.2, 2.2, 4.8, 2.2)
arrow(s, 8.6, 2.2, 9.2, 2.2)
arrow(s, 6.7, 2.75, 6.7, 3.35)
arrow(s, 10.95, 2.75, 10.95, 3.35)
arrow(s, 4.8, 3.9, 4.2, 3.9)
arrow(s, 5.6, 4.45, 2.4, 5.15)
arrow(s, 6.7, 4.45, 6.7, 5.15)
arrow(s, 10.95, 4.45, 10.95, 5.15)

# =====================================================================
# 4. Lab 1
s = new_slide("Spring Boot and the first REST controller", lab=1, variant=7, notes=(
    "Lab 1 set up the environment and the first project. From it the final project keeps the parent POM, the "
    "webmvc starter, DevTools, the Maven Wrapper and HelloController. Individual variant 7 is the wordcount "
    "endpoint: it counts words and characters and finds the longest word. If the parameter is missing, "
    "defaultValue is an empty string, so the answer is zeros instead of a 400 error."))
bullets(s, 0.6, 1.65, 5.6, 3.2, [
    [("Project: ", {"bold": True}), ("Spring Initializr, parent POM 4.1.1, Java 25, Maven Wrapper", {})],
    [("Starters: ", {"bold": True}), ("webmvc (Tomcat and Jackson come transitively), DevTools", {})],
    [("HelloController: ", {"bold": True}), ("@RestController, @GetMapping, @RequestParam, @Value", {})],
    [("Endpoints: ", {"bold": True}), ("/api/hello, /api/info", {})],
    [("Variant 7: ", {"bold": True, "color": GREEN}), ("/api/wordcount — words, characters, longest word", {})],
], size=15)
code(s, 6.7, 1.65, 6.0, 2.85, [
    '@GetMapping("/wordcount")',
    "public WordCount wordCount(",
    '    @RequestParam(defaultValue = "") String text) {',
    "  String trimmed = text.trim();",
    "  if (trimmed.isEmpty())",
    '      return new WordCount(0, 0, "");',
    '  String[] words = trimmed.split("\\\\s+");',
    "  ...",
], size=12, title="web/HelloController.java")
table(s, 0.6, 4.95, 12.1, [
    ["Input text", "wordCount", "characterCount", "longestWord"],
    ["Hello World", "2", "11", "Hello"],
    ["Spring Boot is very interesting", "5", "31", "interesting"],
    ["(parameter omitted)", "0", "0", "\"\""],
], [5.2, 2.2, 2.4, 2.3], size=14, row_h=0.42)

# =====================================================================
# 5. Lab 2
s = new_slide("IoC container: beans, injection, scopes", lab=2, variant=8, notes=(
    "Lab 2 is about the IoC container. There is a Notifier interface with four implementations. Email is marked "
    "Primary, noop is Fallback, and Order defines the order in the list. NotificationService receives them in "
    "every possible way: through Primary, Qualifier, List and Map. Variant 8 is HtmlNotifier, which escapes "
    "angle brackets. In the final project this is not just a demo: when a book is added, the service sends a "
    "notification through the Primary bean. There is also a @Bean for DateTimeFormatter, the bean lifecycle and "
    "a prototype Ticket bean obtained through ObjectProvider."))
order = [("1", "console", "@Order(1)"), ("2", "email", "@Primary"), ("3", "html", "V8 · escape"),
         ("99", "noop", "@Fallback")]
for i, (o, name, tag) in enumerate(order):
    x = 0.6 + i * 1.6
    fill = GREEN if name in ("email", "html") else DARK
    circle(s, x + 0.2, 1.7, 1.1, fill, o, size=22)
    text(s, x, 2.9, 1.5, 0.35, name, size=15, bold=True, color=DARK, align=PP_ALIGN.CENTER)
    text(s, x, 3.25, 1.5, 0.35, tag, size=12, color=MUTED, align=PP_ALIGN.CENTER)
text(s, 0.6, 3.85, 6.2, 0.35, "List<Notifier> is ordered by @Order: console → email → html → noop", size=13,
     color=MUTED)
bullets(s, 0.6, 4.4, 6.2, 2.5, [
    [("@Primary / @Qualifier / List / Map", {"bold": True}), (" — 4 ways to pick a bean", {})],
    [("@Bean reportDateFormatter", {"bold": True}), (" — a JDK class without @Component", {})],
    [("LifecycleDemo:", {"bold": True}), (" constructor → @PostConstruct → @PreDestroy", {})],
    [("Ticket (prototype):", {"bold": True}), (" a new instance only through ObjectProvider", {})],
], size=14, space=6)
rect(s, 7.3, 1.65, 5.4, 2.4, SOFT)
text(s, 7.6, 1.8, 4.8, 0.4, "Variant 8 · HtmlNotifier", size=17, bold=True, color=DARK)
text(s, 7.6, 2.3, 4.8, 1.6, [
    [("input:  ", {"bold": True}), ("Hello <world>", {"font": MONO})],
    [("output: ", {"bold": True}), ("<p>Hello &lt;world&gt;</p>", {"font": MONO})],
    "Angle brackets are escaped instead of being interpreted as HTML",
], size=14, space_after=6)
rect(s, 7.3, 4.35, 5.4, 2.4, DARK)
text(s, 7.6, 4.5, 4.8, 0.4, "How it is used in the project", size=17, bold=True, color=GREEN)
text(s, 7.6, 5.0, 4.8, 1.6,
     "BookService.create() sends «New book in the catalogue» through the @Primary EmailNotifier, "
     "which reads app.mail.* from lab 3", size=14, color=WHITE)

# =====================================================================
# 6. Lab 3
s = new_slide("Configuration: YAML, profiles, properties", lab=3, variant=7, notes=(
    "Lab 3 is about external configuration. application.yml holds the shared values, dev and prod override them. "
    "AppProperties is a record with ConfigurationProperties and Validated: with retry-count set to 99 the "
    "application does not start. EnvironmentBanner shows that for any combination of profiles exactly one bean "
    "exists. Variant 7 is app.cache. In the report it was only printed in JSON; here CitationCache really caches "
    "citations by ttl and max-entries. Precedence: a command-line argument beats a JVM system property, which "
    "beats an environment variable, which beats the files."))
table(s, 0.6, 1.65, 7.3, [
    ["/api/config field", "no profile", "dev", "prod"],
    ["serverPort", "8080", "8081", "8080"],
    ["mailRetryCount", "3", "1", "5"],
    ["mailEnabled", "true", "false", "true"],
    ["cacheTtl (V7)", "PT1M", "PT30S", "PT1M"],
    ["cacheMaxEntries (V7)", "100", "50", "100"],
    ["banner", "no profile", "DEVELOPMENT", "PRODUCTION"],
], [2.5, 1.6, 1.6, 1.6], size=14, row_h=0.45)
text(s, 0.6, 5.0, 7.3, 0.4, "Property source precedence (strongest → weakest)", size=16, bold=True, color=DARK)
steps = ["--arg", "-D JVM", "ENV", "profile.yml", "app.yml"]
for i, st in enumerate(steps):
    x = 0.6 + i * 1.48
    rect(s, x, 5.5, 1.33, 0.7, GREEN if i == 0 else SOFT)
    text(s, x, 5.5, 1.33, 0.7, st, size=13, bold=True, color=WHITE if i == 0 else DARK,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
code(s, 8.4, 1.65, 4.3, 3.0, [
    '@ConfigurationProperties("app")',
    "@Validated",
    "public record AppProperties(",
    "  @NotBlank String owner,",
    "  @Valid Mail mail,",
    "  @Valid Cache cache) {",
    "  record Cache(",
    '   @DefaultValue("1m") Duration ttl,',
    "   @Min(10) @Max(10000)",
    "   int maxEntries) {}",
    "}",
], size=11, title="config/AppProperties.java")
rect(s, 8.4, 4.9, 4.3, 1.75, SOFT)
text(s, 8.65, 5.0, 3.9, 1.6, [
    [("Actuator: ", {"bold": True}), ("/actuator/env/app.mail.retry-count shows the value and its source", {})],
    [("Check: ", {"bold": True}), ("retry-count=99 → APPLICATION FAILED TO START", {})],
], size=13, space_after=6)

# =====================================================================
# 7. Lab 4
s = new_slide("Spring AOP: cross-cutting concerns", lab=4, variant=8, notes=(
    "Lab 4 is AOP. Pointcuts is a library of named pointcuts: public methods of the service package. The aspects "
    "are ordered: audit is the outermost, then logging, then timing. Auditing applies only to methods carrying the "
    "Audited annotation. The service is a CGLIB proxy, which /api/system/proxy shows. Self-invocation: if "
    "deleteBulk calls this.delete, there is no audit; through self there is one record per id. Variant 8 is call "
    "tracing: a citation call, which calls CitationCache.get inside, produces a real nested indentation."))
layers = [(0.6, 1.65, 6.0, 3.9, SOFT, "AuditAspect  @Order(1)  ·  @Around(@annotation)"),
          (0.95, 2.2, 5.3, 3.1, LINE, "LoggingAspect  @Order(2)  ·  @Before / @AfterReturning / @AfterThrowing"),
          (1.3, 2.85, 4.6, 2.2, GREEN, "TimingAspect  @Order(3)  ·  @Around + proceed()"),
          (1.65, 3.45, 3.9, 1.35, DARK, "BookService.delete(id)")]
for x, y, w, h, fill, label in layers:
    rect(s, x, y, w, h, fill)
    text(s, x + 0.2, y + 0.1, w - 0.4, 0.5, label, size=13, bold=True,
         color=WHITE if fill in (GREEN, DARK) else DARK)
text(s, 1.85, 3.95, 3.5, 0.7, "CGLIB proxy:\nBookService$$SpringCGLIB$$1", size=12, color=CODE_TXT, font=MONO)
table(s, 0.6, 5.75, 6.0, [
    ["deleteBulk(ids, viaProxy)", "BOOK_DELETE records"],
    ["false → this.delete()", "0  (proxy bypassed)"],
    ["true → self.delete()", "2  (through the proxy)"],
], [3.4, 2.6], size=12, row_h=0.36)
code(s, 7.1, 1.65, 5.6, 4.0, [
    "[AUDIT] start BOOK_DELETE [3]",
    "[LOG] -> BookService.delete(..) args=[3]",
    "[TRACE]   -> BookService.delete(..)",
    "[TRACE]     -> CitationCache.evict(..)",
    "[TRACE]     <- CitationCache.evict(..)",
    "[TRACE]   <- BookService.delete(..)",
    "[TIME] BookService.delete(..) - 1 ms",
    "[AUDIT] BOOK_DELETE success",
    "",
    "[TIME] SLOW: BookService.statistics() - 301 ms",
], size=11, title="Real log · deleteBulk(viaProxy=true)")
rect(s, 7.1, 5.85, 5.6, 0.95, SOFT)
text(s, 7.35, 5.9, 5.1, 0.85, [[("V8 Call Tracing: ", {"bold": True, "color": GREEN}),
                                 ("@Before increases and @After decreases the ThreadLocal depth — the indent "
                                  "shows the nesting", {})]], size=13, anchor=MSO_ANCHOR.MIDDLE)

# =====================================================================
# 8. Lab 5
s = new_slide("Spring MVC: routing and status codes", lab=5, variant=6, notes=(
    "Lab 5 was done by my partner Gibrat Bekbaev; in the final project it became the base of the catalogue. "
    "A Book domain record, an in-memory repository, a service, a REST controller with five operations and a page "
    "controller with Thymeleaf. The key experiment is the status codes: each of 400, 404, 405, 415 and 406 points "
    "to the part of the mapping that did not match. Variant 6 is a citation as text/plain; with Accept "
    "application/json the answer is 406."))
codes = [("200", "GET /api/books", "list, ?author, ?limit"), ("201", "POST /api/books", "+ Location"),
         ("204", "DELETE /api/books/2", "deleted"), ("404", "DELETE again", "already gone"),
         ("400", "GET /api/books/abc", "abc → long fails"), ("405", "DELETE /api/books", "wrong method"),
         ("415", "POST text/plain", "body media type"), ("406", "Accept: xml", "response media type")]
for i, (c, req, why) in enumerate(codes):
    col, row = i % 4, i // 4
    x, y = 0.6 + col * 3.08, 1.65 + row * 1.65
    ok = c.startswith("2")
    rect(s, x, y, 2.83, 1.45, SOFT if ok else WHITE, line=None if ok else LINE)
    text(s, x + 0.2, y + 0.08, 1.2, 0.7, c, size=32, bold=True, font=HEAD, color=GREEN if ok else RED)
    text(s, x + 0.2, y + 0.78, 2.5, 0.32, req, size=12, bold=True, color=DARK, font=MONO)
    text(s, x + 0.2, y + 1.08, 2.5, 0.3, why, size=12, color=MUTED)
bullets(s, 0.6, 5.1, 6.4, 1.8, [
    "@RestController writes to the response body, @Controller returns a view name",
    "Post/Redirect/Get + flash: no duplicate after F5",
    "static/index.html and css/app.css — served without a controller",
], size=14, space=6)
code(s, 7.25, 5.1, 5.45, 1.7, [
    '@GetMapping(value = "/{id}/citation",',
    '            produces = "text/plain")',
    "→ Joshua Bloch (2018). Effective Java.",
], size=12, title="Variant 6 · citation")

# =====================================================================
# 9. Lab 6
s = new_slide("Forms, validation and error handling", lab=6, variant=7, notes=(
    "Lab 6 is about forms. BookForm is a separate mutable class without an id, which protects from mass "
    "assignment. The path of an error message: the binder fills the form, Valid runs Hibernate Validator, the key "
    "in braces is resolved through the MessageSource for the current locale, the error goes into BindingResult, "
    "the controller returns the form and th:errors prints the text. There is a custom NoDigits constraint and two "
    "ControllerAdvice classes for pages and the API. Variant 7 is the publication date, whose format depends on "
    "the language. In the final project the same constraints also validate the JSON of the REST API."))
flow = [("1", "Browser", "POST /books\ntitle="), ("2", "WebDataBinder", "→ BookForm\ntypeMismatch"),
        ("3", "@Valid", "Hibernate\nValidator"), ("4", "MessageSource", "{book.title.\nrequired} · ru/en"),
        ("5", "BindingResult", "errors, no\nexception"), ("6", "Thymeleaf", "th:errors\nth:errorclass")]
for i, (n, t, sub) in enumerate(flow):
    x = 0.6 + i * 2.05
    circle(s, x + 0.55, 1.65, 0.65, GREEN, n, size=16)
    text(s, x, 2.4, 1.85, 0.35, t, size=14, bold=True, color=DARK, align=PP_ALIGN.CENTER)
    text(s, x, 2.75, 1.85, 0.7, sub, size=12, color=MUTED, align=PP_ALIGN.CENTER)
    if i < 5:
        arrow(s, x + 1.3, 1.975, x + 2.6, 1.975, color=LINE)
rect(s, 0.6, 3.75, 5.9, 3.0, SOFT)
text(s, 0.85, 3.85, 5.4, 0.4, "Variant 7 · publication date", size=17, bold=True, color=DARK)
text(s, 0.85, 4.3, 5.4, 0.9, ["@PastOrPresent", '@DateTimeFormat(style = "S-",', '  fallbackPatterns = "yyyy-MM-dd")'],
     size=12, font=MONO, color=INK)
table(s, 0.85, 5.3, 5.4, [
    ["Locale", "Input", "In the catalogue"],
    ["en", "5/9/06", "9 May 2006"],
    ["ru", "09.05.2006", "9 мая 2006"],
], [1.2, 1.9, 2.3], size=12, row_h=0.4)
bullets(s, 7.0, 3.85, 5.7, 3.0, [
    [("BookForm ", {"bold": True}), ("is separate from the Book record — no id, no mass assignment", {})],
    [("@NoDigits ", {"bold": True}), ("— a custom constraint; null is valid", {})],
    [("messages(_ru).properties ", {"bold": True}), ("+ ?lang=ru via SessionLocaleResolver", {})],
    [("/books/999 ", {"bold": True}), ("→ HTML 404 page (PageExceptionHandler)", {})],
    [("/api/books/999 ", {"bold": True}), ("→ application/problem+json (ApiExceptionHandler)", {})],
], size=14, space=6)

# =====================================================================
# 10. Integration
s = new_slide("Not side by side, but together: how labs connect", notes=(
    "The main difference between the final project and a set of labs is the connections. Configuration controls "
    "notifications and the cache. Aspects wrap every catalogue operation. Validation from lab 6 also protects the "
    "REST API, and the IllegalArgumentException from lab 4 now becomes 400 instead of 500."))
links = [("2 → 5", "New book notifications", "BookService.create() calls the @Primary EmailNotifier"),
         ("3 → 2", "Mail depends on the profile", "dev: mail disabled · prod: retry-count = 5"),
         ("3 → 5", "Citation cache from app.cache", "1 min / 100 entries, under dev — 30 s / 50"),
         ("4 → all", "Every operation is wrapped", "Audit · Log · Time · Trace on the whole service layer"),
         ("6 → 5", "JSON validation in the API", "BookRequest: the same constraints and messages → 400"),
         ("4 + 6", "Client error is 400", "delete(0): @AfterThrowing, audit failure, ProblemDetail")]
for i, (tag, head, body) in enumerate(links):
    col, row = i % 3, i // 3
    x, y = 0.6 + col * 4.08, 1.65 + row * 2.6
    rect(s, x, y, 3.83, 2.35, SOFT)
    pill = rect(s, x + 0.25, y + 0.25, 1.35, 0.45, DARK)
    pill.adjustments[0] = 0.5
    text(s, x + 0.25, y + 0.25, 1.35, 0.45, tag, size=13, bold=True, color=GREEN,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.25, y + 0.9, 3.35, 0.4, head, size=16, bold=True, color=DARK)
    text(s, x + 0.25, y + 1.35, 3.35, 0.9, body, size=13, color=MUTED)

# =====================================================================
# 11. Tests (native chart)
s = new_slide("Every experiment is an automated test", notes=(
    "So that the defence does not depend on screenshots, every experiment from the reports is written as a "
    "MockMvc integration test. There are 17 of them: they check status codes, validation errors, Russian "
    "messages, the proxy, self-invocation and the dev profile. One command, mvnw test, and everything is green."))
cd = CategoryChartData()
cd.categories = ["Lab 1", "Lab 2", "Lab 3", "Lab 4", "Lab 5", "Lab 6"]
cd.add_series("Tests", (1, 2, 2, 3, 4, 5))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.6), Inches(7.2), Inches(5.2), cd)
ch = gf.chart
ch.has_legend = False
ch.has_title = True
ch.chart_title.text_frame.text = "Tests per lab"
ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
ch.chart_title.text_frame.paragraphs[0].runs[0].font.color.rgb = MUTED
plot = ch.plots[0]
plot.gap_width = 60
plot.has_data_labels = True
plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
plot.data_labels.font.size = Pt(14)
plot.data_labels.font.bold = True
plot.data_labels.font.color.rgb = DARK
ser = plot.series[0]
ser.format.fill.solid()
ser.format.fill.fore_color.rgb = GREEN
ch.value_axis.visible = False
ch.value_axis.has_major_gridlines = False
ch.category_axis.tick_labels.font.size = Pt(13)
ch.category_axis.tick_labels.font.color.rgb = MUTED
ch.category_axis.format.line.color.rgb = LINE
rect(s, 8.3, 1.65, 4.4, 2.0, DARK)
text(s, 8.3, 1.75, 4.4, 1.0, "17 / 17", size=48, bold=True, font=HEAD, color=GREEN, align=PP_ALIGN.CENTER)
text(s, 8.3, 2.85, 4.4, 0.5, "./mvnw test  →  BUILD SUCCESS", size=14, color=WHITE, font=MONO,
     align=PP_ALIGN.CENTER)
bullets(s, 8.3, 4.0, 4.4, 2.9, [
    "200 / 201 / 204 / 400 / 404 / 405 / 406 / 415",
    "form errors in EN and RU, flash + redirect",
    "CGLIB proxy and self-invocation: 0 vs 2 records",
    "dev profile: port 8081, ttl PT30S",
], size=14, space=6)

# =====================================================================
# 12. Demo
s = new_slide("Demonstration script", notes=(
    "The demonstration takes about five minutes. First the tests, then the start page with a card for each lab, "
    "the Russian form with errors, adding a book correctly and the aspect log in the console, the "
    "self-invocation experiment, two different 404 responses and a restart with the dev profile."))
demo = [("./mvnw test", "17 tests are green"),
        ("./mvnw spring-boot:run → /", "start page: links for every lab"),
        ("/books/new?lang=ru → «Сохранить»", "all validation errors in Russian"),
        ("Valid form → F5", "redirect + flash, no duplicate; [AUDIT] and EMAIL in the log"),
        ("DELETE /api/books/bulk?ids=1,2&viaProxy=false|true", "/api/system/audit: 0 and 2 records"),
        ("/books/999 and /api/books/999", "HTML page vs ProblemDetail"),
        ("-Dspring-boot.run.profiles=dev → /api/config", "port 8081, PT30S, DEVELOPMENT")]
for i, (cmd, res) in enumerate(demo):
    y = 1.6 + i * 0.75
    circle(s, 0.6, y + 0.05, 0.5, GREEN if i % 2 == 0 else DARK, str(i + 1), size=14)
    text(s, 1.35, y, 6.6, 0.6, cmd, size=14, bold=True, font=MONO, color=DARK, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 8.0, y, 4.7, 0.6, res, size=14, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)

# =====================================================================
# 13. Q&A
s = new_slide("Key defence questions", notes=(
    "These are the most frequent questions across all labs. The full list with answers is in docs/DEFENSE.md."))
qa = [("Why constructor injection?", "final fields, the object is always complete, easy to test"),
      ("A prototype inside a singleton?", "created once; a fresh one only through ObjectProvider"),
      ("Profile file or --argument?", "the command-line argument: a profile is just a PropertySource"),
      ("What if proceed() is removed?", "the target method is never executed"),
      ("Why is self-invocation not advised?", "this.method() skips the proxy → use a self reference"),
      ("What do 405 / 415 / 406 mean?", "wrong method / body media type / response media type"),
      ("Why a separate BookForm?", "protects from mass assignment, keeps the user input"),
      ("BindingResult not right after the form?", "an exception and 400 instead of the form with errors")]
for i, (q, a) in enumerate(qa):
    col, row = i % 2, i // 2
    x, y = 0.6 + col * 6.15, 1.6 + row * 1.33
    rect(s, x, y, 5.9, 1.15, SOFT if (row + col) % 2 == 0 else WHITE, line=None if (row + col) % 2 == 0 else LINE)
    text(s, x + 0.25, y + 0.12, 5.4, 0.4, q, size=15, bold=True, color=DARK)
    text(s, x + 0.25, y + 0.55, 5.4, 0.5, a, size=13, color=MUTED)

# =====================================================================
# 14. Closing
s = new_slide(dark=True, notes=(
    "To sum up: six laboratory works are combined into one working application where each topic is a layer of the "
    "system, the individual variants are implemented and connected to each other, and the experiments are "
    "verified by tests. Thank you for your attention, I am ready to answer your questions."))
text(s, 0.9, 1.0, 11.5, 1.0, "Summary", size=48, bold=True, font=HEAD, color=WHITE)
results = [("One application", "a library catalogue instead of six separate demos"),
           ("Six variants", "V7 · V8 · V7 · V8 · V6 · V7 — implemented and connected"),
           ("17 tests", "every experiment from the reports is reproduced")]
for i, (h, b) in enumerate(results):
    x = 0.9 + i * 3.95
    circle(s, x, 2.5, 0.75, GREEN, str(i + 1), size=20)
    text(s, x, 3.45, 3.6, 0.5, h, size=22, bold=True, color=WHITE)
    text(s, x, 4.0, 3.6, 1.0, b, size=15, color=CODE_TXT)
text(s, 0.9, 5.6, 11.5, 0.6, "Thank you for your attention! Questions are welcome.", size=24, color=WHITE)
text(s, 0.9, 6.3, 11.5, 0.4, REPO, size=16, color=GREEN)

out = Path(__file__).with_name("Defense_Presentation.pptx")
prs.save(out)
print("saved", out, "slides:", len(prs.slides))
